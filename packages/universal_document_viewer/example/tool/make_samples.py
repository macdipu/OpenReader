"""Builds the example app's sample documents in assets/samples.

    pip install reportlab openpyxl
    python3 tool/make_samples.py

The Word sample is the kit's own test document
(../test/fixtures/word/features.docx, from its generate.py) and the
PowerPoint sample is built with the kit's `buildPptx` (see README.md).
"""
import os
import shutil

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, Rule
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from reportlab.lib import pdfencrypt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'assets', 'samples')
WORDS = ('alpha beta gamma delta epsilon zeta eta theta iota kappa lambda mu nu xi omicron pi rho sigma tau '
         'upsilon phi chi psi omega').split()


def sentence(n, length=14):
    return ' '.join(WORDS[(n * 7 + i * 3) % len(WORDS)] for i in range(length)).capitalize() + '.'


def guide_pdf(name='guide.pdf', password=None):
    """Twelve A4 pages with a heading each, so page and resume are easy to see."""
    encrypt = pdfencrypt.StandardEncryption(password, canPrint=1) if password else None
    c = canvas.Canvas(os.path.join(OUT, name), pagesize=A4, encrypt=encrypt)
    c.setTitle('universal_document_viewer sample guide')
    for page in range(1, 13):
        c.setFont('Helvetica-Bold', 22)
        c.drawString(56, 780, f'Chapter {page}')
        c.setFont('Helvetica', 11)
        y = 750
        for line in range(42):
            c.drawString(56, y, sentence(page * 100 + line))
            y -= 16
        if page == 1:
            c.setFillColorRGB(0.1, 0.3, 0.8)
            c.drawString(56, 60, 'Flutter website (a web link)')
            c.linkURL('https://flutter.dev/', (56, 55, 230, 72), relative=0)
            c.setFillColorRGB(0, 0, 0)
        c.setFont('Helvetica', 9)
        c.drawString(500, 30, f'{page} / 12')
        c.showPage()
    c.save()


def sales_xlsx():
    """Borders, conditional formatting, notes, formulas and a long sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Regions'
    thin, thick, double = Side(style='thin'), Side(style='thick'), Side(style='double')
    ws.append(['Region', 'Units', 'Target', 'Status', 'Score'])
    rows = [('North', 120, 100, 'On track', 35), ('South', 80, 100, 'Overdue', 90), ('East', 150, 120, 'On track', 60),
            ('West', 40, 90, 'Overdue', 10), ('Central', 95, 95, 'Review', 75)]
    for row in rows:
        ws.append(list(row))
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='DDEBF7')
        cell.border = Border(bottom=thick, top=thin, left=thin, right=thin)
    for row in ws.iter_rows(min_row=2, max_row=6, max_col=5):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws['A8'] = 'Total'
    ws['A8'].font = Font(bold=True)
    ws['B8'] = '=SUM(B2:B6)'
    ws['C8'] = '=SUM(C2:C6)'
    for cell in ws[8][:5]:
        cell.border = Border(top=thin, bottom=double)
    red = DifferentialStyle(font=Font(color='9C0006', bold=True), fill=PatternFill(bgColor='FFC7CE'))
    ws.conditional_formatting.add('B2:B6', CellIsRule(operator='lessThan', formula=['C2'], font=Font(color='9C0006'),
                                                      fill=PatternFill(bgColor='FFC7CE')))
    overdue = Rule(type='containsText', operator='containsText', text='Overdue', dxf=red)
    overdue.formula = ['NOT(ISERROR(SEARCH("Overdue",D2)))']
    ws.conditional_formatting.add('D2:D6', overdue)
    ws.conditional_formatting.add('E2:E6', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=100,
                                                       color='638EC6'))
    ws['B3'].comment = Comment('Shipment delayed - confirm with logistics', 'Rina')
    ws['A8'].comment = Comment('Excludes returns', 'Sam')
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['D'].width = 12

    orders = wb.create_sheet('Orders')
    orders.append(['Order', 'Customer', 'Amount', 'Status'])
    for cell in orders[1]:
        cell.font = Font(bold=True)
    for i in range(1, 301):
        orders.append([f'ORD-{i:04d}', f'Customer {i % 37}', i * 13 % 997, 'Open' if i % 5 else 'Late'])
    orders.conditional_formatting.add('C2:C301', ColorScaleRule(start_type='min', start_color='F8696B',
                                                                mid_type='percentile', mid_value=50,
                                                                mid_color='FFEB84', end_type='max',
                                                                end_color='63BE7B'))
    orders.column_dimensions['B'].width = 14
    wb.save(os.path.join(OUT, 'sales.xlsx'))


def data_csv():
    with open(os.path.join(OUT, 'contacts.csv'), 'w', encoding='utf-8', newline='') as f:
        f.write('Name,City,Phone,Note\n')
        cities = ['Dhaka', 'Chattogram', 'Sylhet', 'Khulna', 'Rajshahi']
        for i in range(1, 121):
            f.write(f'Person {i},{cities[i % 5]},+880 1{i:09d},"Met at event {i % 7}, follow up"\n')


def notes_txt():
    with open(os.path.join(OUT, 'notes.txt'), 'w', encoding='utf-8') as f:
        for i in range(1, 601):
            f.write(f'{i:>4}  {sentence(i, 6 + i % 18)}\n')


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    guide_pdf()
    guide_pdf('locked.pdf', password='open')
    sales_xlsx()
    data_csv()
    notes_txt()
    shutil.copy(os.path.join(HERE, '..', '..', 'test', 'fixtures', 'word', 'features.docx'), os.path.join(OUT, 'report.docx'))
    print('Samples written to', os.path.normpath(OUT))
