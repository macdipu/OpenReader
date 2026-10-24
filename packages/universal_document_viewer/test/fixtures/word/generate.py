"""Builds the .docx fixtures for the Word engine tests with python-docx, an
independent writer (so the reader isn't only tested against files it could
have written itself).

    python3 test/fixtures/word/generate.py
"""
import os
import struct
import zlib

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

HERE = os.path.dirname(os.path.abspath(__file__))


def png(width, height, rgb):
    """A solid-colour PNG, without an imaging library."""
    raw = b''.join(b'\x00' + bytes(rgb) * width for _ in range(height))
    def chunk(kind, data):
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data) & 0xffffffff)
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', width, height, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b''))


def hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                         is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    style = OxmlElement('w:rStyle')
    style.set(qn('w:val'), 'Hyperlink')
    rpr.append(style)
    run.append(rpr)
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def features():
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = 'Report header'
    section.footer.paragraphs[0].text = 'Page footer'

    doc.add_heading('Quarterly Report', level=1)
    p = doc.add_paragraph('Plain, ')
    p.add_run('bold').bold = True
    p.add_run(', ')
    p.add_run('italic').italic = True
    p.add_run(', ')
    u = p.add_run('red underline')
    u.underline = True
    u.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    u.font.size = Pt(14)
    p.add_run('. Visit ')
    hyperlink(p, 'our site', 'https://example.com/')
    p.add_run('.')

    centred = doc.add_paragraph('Centred line')
    centred.alignment = WD_ALIGN_PARAGRAPH.CENTER

    tabs = doc.add_paragraph()
    run = tabs.add_run('left\tright')
    run.add_break()
    run.add_text('second line')

    for item in ['Apples', 'Pears']:
        doc.add_paragraph(item, style='List Bullet')
    for item in ['First', 'Second', 'Third']:
        doc.add_paragraph(item, style='List Number')

    table = doc.add_table(rows=3, cols=3)
    table.style = 'Table Grid'
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f'r{r}c{c}'
    table.cell(0, 0).merge(table.cell(0, 1))     # spans two columns
    table.cell(1, 2).merge(table.cell(2, 2))     # spans two rows
    table.cell(0, 0).text = 'wide'
    table.cell(1, 2).text = 'tall'

    doc.add_paragraph().add_run().add_picture(_image_file(), width=Inches(1), height=Inches(0.5))

    doc.add_paragraph('Before the break').add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph('After the break')

    landscape = doc.add_section()
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = section.page_height, section.page_width
    doc.add_paragraph('Landscape section')

    doc.save(os.path.join(HERE, 'features.docx'))


def _image_file():
    path = os.path.join(HERE, '_tmp.png')
    with open(path, 'wb') as f:
        f.write(png(4, 2, (0x44, 0x72, 0xC4)))
    return path


if __name__ == '__main__':
    features()
    os.remove(os.path.join(HERE, '_tmp.png'))
