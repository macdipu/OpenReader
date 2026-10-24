"""Builds the .xlsx fixtures for the sheet engine tests with openpyxl, an
independent writer (so the reader isn't only tested against files it could
have written itself).

    python3 test/fixtures/sheets/generate.py
"""
import datetime
import os

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.colors import Color

HERE = os.path.dirname(os.path.abspath(__file__))


def features():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = 'Name'
    ws['B1'] = 'Score'
    ws['C1'] = 'Joined'
    ws['D1'] = 'Share'
    ws['E1'] = 'Paid'
    ws['F1'] = 'Active'
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws['A2'] = 'Alice'
    ws['B2'] = 97
    ws['C2'] = datetime.date(2024, 1, 15)
    ws['C2'].number_format = 'yyyy-mm-dd'
    ws['D2'] = 0.256
    ws['D2'].number_format = '0.0%'
    ws['E2'] = 1234.5
    ws['E2'].number_format = '#,##0.00'
    ws['F2'] = True
    ws['A3'] = 'Bob'
    ws['B3'] = 81.25
    ws['C3'] = datetime.datetime(2023, 7, 4, 18, 30)
    ws['C3'].number_format = 'm/d/yyyy h:mm'
    ws['D3'] = -0.05
    ws['D3'].number_format = '0%'
    ws['E3'] = -42
    ws['E3'].number_format = '$#,##0.00_);[Red]($#,##0.00)'
    ws['F3'] = False
    ws['A4'] = '=SUM(B2:B3)'          # formula saved without a result
    ws['A5'] = CellRichText([TextBlock(InlineFont(b=True), 'Rich '), 'text'])
    ws['A6'] = 'Italic'
    ws['A6'].font = Font(italic=True)
    ws['B6'] = 'rgb fill'
    ws['B6'].fill = PatternFill('solid', fgColor='FFC000')
    ws['C6'] = 'theme fill'
    ws['C6'].fill = PatternFill('solid', fgColor=Color(theme=4, tint=0.3999755851924192))
    ws['D6'] = 'indexed fill'
    ws['D6'].fill = PatternFill('solid', fgColor=Color(indexed=10))
    ws['E6'].fill = PatternFill('solid', fgColor='92D050')   # blank but filled
    ws['A7'] = '#N/A'
    ws['A7'].data_type = 'e'
    ws['A10'] = 'merged'
    ws.merge_cells('A10:C10')      # crosses the frozen column: left unmerged
    ws['B13'] = 'wide'
    ws.merge_cells('B13:D13')
    ws['H12'] = 'far corner'
    ws.freeze_panes = 'B2'

    # Layout: widths, heights, hidden rows/columns.
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['G'].hidden = True
    ws.row_dimensions[8].height = 30
    ws.row_dimensions[9].hidden = True
    # Text formatting and alignment.
    ws['B7'] = 'red underline'
    ws['B7'].font = Font(color='FFC00000', underline='single', size=14)
    ws['C7'] = 'struck'
    ws['C7'].font = Font(strike=True)
    ws['D7'] = 'centred'
    ws['D7'].alignment = Alignment(horizontal='center', vertical='top')
    ws['E7'] = 'wrapped long text'
    ws['E7'].alignment = Alignment(wrap_text=True)
    # Links: a web address and a place in the workbook.
    ws['B11'] = 'site'
    ws['B11'].hyperlink = 'https://example.com/'
    ws['D11'] = 'jump'
    ws['D11'].hyperlink = '#Later!B3'

    secret = wb.create_sheet('Secret')
    secret['A1'] = 'hidden'
    secret.sheet_state = 'hidden'
    very = wb.create_sheet('VeryHidden')
    very['A1'] = 'very hidden'
    very.sheet_state = 'veryHidden'
    later = wb.create_sheet('Later')
    later['B3'] = 'alice again'
    wb.save(os.path.join(HERE, 'features.xlsx'))


def large():
    wb = Workbook(write_only=True)
    ws = wb.create_sheet('Data')
    ws.append(['Id', 'Name', 'City', 'Amount', 'Date', 'Rate', 'Code', 'Note'])
    start = datetime.date(2020, 1, 1)
    for i in range(1, 6440):
        ws.append([i, f'Person {i}', ['Dhaka', 'Paris', 'Lima'][i % 3], i * 1.5,
                   start + datetime.timedelta(days=i), (i % 100) / 100, f'C{i:05d}', 'x' * (i % 20)])
    wb.save(os.path.join(HERE, 'large.xlsx'))


if __name__ == '__main__':
    features()
    large()
